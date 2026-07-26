"""
Generate README_SIMM.docx — dokumentasi proyek SIMM dalam format Word.
Jalankan: python generate_readme_docx.py
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ── Page margins (A4) ────────────────────────────────────────────────────────
section = doc.sections[0]
section.page_width    = Inches(8.27)
section.page_height   = Inches(11.69)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.5)
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)

# ── Base style ────────────────────────────────────────────────────────────────
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)

# ── Helpers ───────────────────────────────────────────────────────────────────

BLUE_DARK  = RGBColor(0x1e, 0x40, 0xaf)
BLUE_MID   = RGBColor(0x1d, 0x4e, 0x89)
BLUE_LIGHT = RGBColor(0x15, 0x65, 0xc0)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT  = RGBColor(0x1e, 0x29, 0x3b)


def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    colors = {1: BLUE_DARK, 2: BLUE_MID, 3: BLUE_LIGHT}
    sizes  = {1: 18, 2: 14, 3: 12}
    for run in h.runs:
        run.font.name  = "Calibri"
        run.font.size  = Pt(sizes.get(level, 11))
        run.font.color.rgb = colors.get(level, DARK_TEXT)
    return h


def para(text, bold=False, italic=False, size=11, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold       = bold
    run.italic     = italic
    run.font.size  = Pt(size)
    run.font.name  = "Calibri"
    return p


def code_block(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  "F1F5F9")
    pPr.append(shd)
    run = p.add_run(text)
    run.font.name      = "Courier New"
    run.font.size      = Pt(9)
    run.font.color.rgb = DARK_TEXT
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    return p


def table(headers, rows, header_color="1E40AF", alt_row_color="EFF6FF"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style     = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header row
    hdr_row = t.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = h
        run = cell.paragraphs[0].runs[0]
        run.bold            = True
        run.font.name       = "Calibri"
        run.font.size       = Pt(10)
        run.font.color.rgb  = WHITE
        tcPr = cell._tc.get_or_add_tcPr()
        shd  = OxmlElement("w:shd")
        shd.set(qn("w:val"),   "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"),  header_color)
        tcPr.append(shd)

    # Data rows
    for ri, row_data in enumerate(rows):
        tr = t.rows[ri + 1]
        for ci, val in enumerate(row_data):
            cell = tr.cells[ci]
            cell.text = val
            run = cell.paragraphs[0].runs[0]
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            if ri % 2 == 1:
                tcPr = cell._tc.get_or_add_tcPr()
                shd  = OxmlElement("w:shd")
                shd.set(qn("w:val"),   "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"),  alt_row_color)
                tcPr.append(shd)
    return t


def gap():
    doc.add_paragraph()


# ============================================================
#  HALAMAN JUDUL
# ============================================================
gap()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SIMM")
r.bold, r.font.size, r.font.name = True, Pt(36), "Calibri"
r.font.color.rgb = BLUE_DARK

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Sistem Informasi Manajemen Maintenance Motor")
r.bold, r.font.size, r.font.name = True, Pt(18), "Calibri"
r.font.color.rgb = BLUE_MID

gap()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Aplikasi web berbasis Laravel 11 untuk manajemen preventive maintenance "
    "motor listrik industri, dilengkapi fitur keamanan data (Data Protection) "
    "berupa Password Hashing, Digital Signature, Activity Log, serta fitur "
    "Backup & Restore Database."
)
r.italic, r.font.size, r.font.name = True, Pt(12), "Calibri"

gap(); gap()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Dikembangkan untuk keperluan Ujian Akhir Semester (UAS)\n"
    "Keamanan Data dan Informasi"
)
r.font.size, r.font.name = Pt(11), "Calibri"

doc.add_page_break()

# ============================================================
#  TECH STACK
# ============================================================
heading("Tech Stack", 1)
table(
    ["Komponen", "Teknologi"],
    [
        ["Backend",      "PHP 8.2, Laravel 11"],
        ["Database",     "MySQL 8"],
        ["Frontend",     "Bootstrap 5, Feather Icons, Viho Admin Template"],
        ["WebSocket",    "Laravel Reverb"],
        ["PDF Export",   "barryvdh/laravel-dompdf"],
        ["Excel Export", "maatwebsite/laravel-excel"],
        ["Queue",        "Database Queue"],
        ["Auth",         "Session-based (tanpa Breeze/Jetstream)"],
    ],
)
gap()

# ============================================================
#  FITUR UTAMA
# ============================================================
heading("Fitur Utama", 1)

FEATURES = {
    "Dashboard": [
        "Stat cards: total motor, jadwal pending, overdue, dan maintenance bulan ini",
        "Tabel jadwal mendekati deadline (H-3)",
        "Real-time toast notification via Laravel Reverb WebSocket",
    ],
    "Motors": [
        "CRUD lengkap motor industri (kode motor, lokasi, area, kategori HP)",
        "Soft delete dengan kemampuan restore",
        "Riwayat maintenance per motor",
    ],
    "Schedules": [
        "CRUD jadwal maintenance",
        "Filter status (pending/overdue/done)",
        "Auto-deteksi overdue via Artisan scheduler",
    ],
    "Maintenance Input": [
        "Pilih jadwal -> data motor otomatis terisi",
        "Checklist 15 aktivitas maintenance",
        "Upload foto bukti inspeksi",
        "Digital Signature otomatis di-generate saat log disimpan",
    ],
    "Reports": [
        "Filter berdasarkan bulan dan motor",
        "Export ke PDF (DomPDF) dan Excel (Maatwebsite)",
    ],
    "Activity Log": [
        "Rekam jejak seluruh aktivitas admin (login, logout, CRUD, dll)",
        "Deteksi dan pencatatan akses ilegal / tampering data",
        "Filter berdasarkan status, modul, dan keyword",
    ],
    "Backup & Restore": [
        "Backup database ke file .sql tanpa bergantung pada mysqldump",
        "Backup file/aset storage ke file .zip",
        "Download dan hapus file backup",
        "Restore database dari file .sql yang diupload maupun dari backup yang tersimpan",
    ],
}

for feat, items in FEATURES.items():
    heading(feat, 2)
    for item in items:
        bullet(item)

gap()

# ============================================================
#  FITUR KEAMANAN
# ============================================================
heading("Fitur Keamanan (Data Protection)", 1)

# 1. Password Hashing
heading("1. Password Hashing", 2)
para(
    "Semua password admin di-hash menggunakan algoritma Bcrypt sebelum disimpan ke "
    "database. Tidak ada password yang pernah tersimpan dalam bentuk plaintext."
)
gap()
para("Implementasi:", bold=True)
code_block(
    "// app/Models/User.php\n"
    "protected function casts(): array {\n"
    "    return ['password' => 'hashed'];\n"
    "}\n\n"
    "// AuthController.php -- Register\n"
    "'password' => Hash::make($validated['password'])\n\n"
    "// Login (verifikasi hash otomatis)\n"
    "Auth::attempt(['username' => ..., 'password' => ...])"
)
gap()
para("Konfigurasi Bcrypt rounds di .env:", bold=True)
code_block("BCRYPT_ROUNDS=12")
gap()

# 2. CAPTCHA
heading("2. CAPTCHA pada Login", 2)
para(
    "Setiap percobaan login dilindungi oleh CAPTCHA SVG dinamis yang di-generate "
    "server-side (tanpa library pihak ketiga). CAPTCHA terdiri dari 5 karakter "
    "alphanumeric acak dengan distorsi, noise line, dan noise dot untuk mencegah "
    "bot otomatis."
)
gap()
para("Alur validasi:", bold=True)
bullet("User input CAPTCHA dibandingkan dengan session captcha_code")
bullet("GAGAL: ActivityLog dicatat (status: warning) + CAPTCHA di-refresh")
bullet("BERHASIL: lanjut verifikasi username/password")
gap()

# 3. Rate Limiting
heading("3. Rate Limiting Login", 2)
para(
    "Sistem membatasi percobaan login sebanyak 5 kali per kombinasi username + "
    "IP address. Jika batas terlampaui, akun dikunci sementara selama 5 menit."
)
gap()
table(
    ["Kondisi", "Respon Sistem"],
    [
        ["CAPTCHA salah",              "Error + log warning"],
        ["Password salah (1-4x)",      "Error + sisa percobaan ditampilkan"],
        ["Password salah (5x)",        "Akun terkunci 5 menit + log danger"],
        ["Akun terkunci, akses ulang", "Tampilkan sisa waktu tunggu"],
    ],
)
gap()

# 4. Digital Signature
heading("4. Digital Signature", 2)
para(
    "Setiap maintenance log yang disimpan akan mendapatkan Digital Signature "
    "(HMAC-SHA256) yang di-generate dari field-field krusial. Jika data diubah "
    "secara ilegal, sistem akan mendeteksi pelanggaran integritas data secara "
    "otomatis."
)
gap()
para("Field yang diproteksi per log:", bold=True)
code_block("admin_id | motor_id | schedule_id | inspection_date | general_notes")
gap()
para("Cara kerja sistem:", bold=True)
code_block(
    "[Admin simpan Maintenance Log]\n"
    "    -> generateSignature() -> HMAC-SHA256 dari payload\n"
    "    -> digital_signature + payload_snapshot disimpan ke DB\n\n"
    "Setiap 30 detik:\n"
    "Browser -> GET /integrity-check -> verifySignature()\n"
    "    -> MISMATCH -> getDiff() -> Banner merah + Activity Log Danger"
)
gap()
para("Before/After Diff saat tampering terdeteksi:", bold=True)
table(
    ["Field", "Sebelum (Asli)", "Sesudah (Diubah)"],
    [["Catatan Umum", "Motor normal, tidak ada kerusakan", "DATA DIUBAH VIA PHPMYADMIN"]],
)
gap()
para("Vektor illegal yang terdeteksi:", bold=True)
for v in [
    "phpMyAdmin / MySQL Workbench / DBeaver / HeidiSQL / TablePlus",
    "MySQL CLI (mysql -u root -p)",
    "php artisan tinker dengan raw DB::table()->update()",
    "SQL Injection (jika ada celah di form)",
    "SSH direct query ke database",
    "Script PHP lain di server yang bypass Laravel",
    "Restore backup lama yang menimpa data",
]:
    bullet(v)
gap()

# 5. Activity Log
heading("5. Activity Log", 2)
para(
    "Seluruh aktivitas admin dicatat otomatis ke tabel activity_logs dengan informasi: "
    "waktu, modul, aksi, deskripsi, status keparahan, IP address, dan user agent."
)
gap()
table(
    ["Aksi", "Modul", "Status"],
    [
        ["Login berhasil",               "Auth",        "Normal"],
        ["Login gagal - CAPTCHA salah",  "Auth",        "Warning"],
        ["Login gagal - password salah", "Auth",        "Warning"],
        ["Login terkunci (rate limit)",  "Auth",        "Danger"],
        ["Register admin baru",          "Auth",        "Normal"],
        ["Logout",                       "Auth",        "Normal"],
        ["Buat maintenance log",         "Maintenance", "Normal"],
        ["Hapus maintenance log",        "Maintenance", "Warning"],
        ["Backup database dibuat",       "Backup",      "Normal"],
        ["Backup dihapus",               "Backup",      "Warning"],
        ["Restore database",             "Backup",      "Warning"],
        ["Tampering terdeteksi",         "Integrity",   "Danger"],
    ],
)
gap()

# ============================================================
#  FITUR BACKUP & RESTORE
# ============================================================
heading("Fitur Backup & Restore", 1)

heading("Backup Database (.sql)", 2)
para(
    "Backup dilakukan via PHP PDO — tidak membutuhkan binary mysqldump di server. "
    "Semua tabel dan data di-dump ke file .sql yang mencakup perintah DROP TABLE IF EXISTS, "
    "CREATE TABLE, dan INSERT INTO. Foreign Key Checks dinonaktifkan sementara saat dump."
)
gap()

heading("Backup Files/Storage (.zip)", 2)
para(
    "Seluruh isi direktori storage/app/public (termasuk foto bukti inspeksi) dikompres "
    "ke file .zip menggunakan ekstensi ZipArchive PHP."
)
gap()

heading("Restore", 2)
table(
    ["Metode", "Keterangan"],
    [
        ["Upload file .sql",      "Restore dari komputer lokal (maks 100 MB)"],
        ["Restore dari storage",  "Restore langsung dari file backup yang tersimpan di server"],
    ],
)
gap()
para(
    "Perhatian: Restore akan menimpa seluruh data database. Setelah restore, semua "
    "Digital Signature mungkin perlu di-regenerasi menggunakan perintah "
    "php artisan maintenance:backfill-signatures apabila APP_KEY telah berubah.",
    italic=True,
)
gap()

# ============================================================
#  STRUKTUR DATABASE
# ============================================================
heading("Struktur Database", 1)
table(
    ["Tabel", "Keterangan"],
    [
        ["users",                        "Akun admin & teknisi"],
        ["motors",                       "Data motor industri (soft delete)"],
        ["schedules",                    "Jadwal maintenance per motor"],
        ["activities",                   "Master daftar aktivitas checklist"],
        ["maintenance_logs",             "Log inspeksi (+ digital_signature + payload_snapshot)"],
        ["maintenance_activity_details", "Detail checklist per log"],
        ["activity_logs",                "Rekam jejak semua aktivitas & kejadian keamanan"],
    ],
)
gap()

# ============================================================
#  ROLE & HAK AKSES
# ============================================================
heading("Role & Hak Akses", 1)
table(
    ["Fitur", "Admin", "Teknisi"],
    [
        ["Dashboard",        "Ya",    "Tidak"],
        ["CRUD Motor",       "Ya",    "Tidak"],
        ["CRUD Jadwal",      "Ya",    "Tidak"],
        ["Input Maintenance","Ya",    "Ya"],
        ["Laporan & Export", "Ya",    "Tidak"],
        ["Activity Log",     "Ya",    "Tidak"],
        ["Backup & Restore", "Ya",    "Tidak"],
        ["Manajemen Profil", "Ya",    "Ya"],
    ],
)
gap()

# ============================================================
#  INSTALASI
# ============================================================
heading("Instalasi", 1)
heading("Prasyarat", 2)
for p_text in ["PHP >= 8.2 (dengan ekstensi: zip, pdo_mysql, gd)", "Composer", "Node.js & NPM", "MySQL 8"]:
    bullet(p_text)
gap()

heading("Langkah-langkah", 2)
steps = [
    ("1. Clone repository",
     "git clone https://github.com/username/SIMM.git\ncd SIMM"),
    ("2. Install dependensi PHP",
     "composer install"),
    ("3. Install dependensi Node",
     "npm install"),
    ("4. Salin file environment",
     "cp .env.example .env"),
    ("5. Generate application key",
     "php artisan key:generate"),
    ("6. Buat database MySQL",
     "CREATE DATABASE motor_pm CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;"),
    ("7. Jalankan migrasi",
     "php artisan migrate"),
    ("8. Seed database (data awal)",
     "php artisan db:seed"),
    ("9. Buat storage symlink",
     "php artisan storage:link"),
    ("10. Build frontend assets",
     "npm run dev"),
    ("11. Generate Digital Signature untuk data yang sudah ada",
     "php artisan maintenance:backfill-signatures"),
]
for title, cmd in steps:
    para(title, bold=True)
    code_block(cmd)
    gap()

# ============================================================
#  KONFIGURASI ENVIRONMENT
# ============================================================
heading("Konfigurasi Environment", 1)
code_block(
    'APP_NAME="SIMM - Motor PM System"\n'
    "APP_ENV=production\n"
    "APP_KEY=                     # Di-generate via key:generate\n"
    "APP_DEBUG=false\n"
    "APP_URL=http://your-domain.com\n\n"
    "DB_CONNECTION=mysql\n"
    "DB_HOST=127.0.0.1\n"
    "DB_PORT=3306\n"
    "DB_DATABASE=motor_pm\n"
    "DB_USERNAME=root\n"
    "DB_PASSWORD=your_password\n\n"
    "QUEUE_CONNECTION=database\n\n"
    "BROADCAST_CONNECTION=reverb\n"
    "REVERB_APP_ID=motor-pm-app\n"
    "REVERB_APP_KEY=your-reverb-key\n"
    "REVERB_APP_SECRET=your-reverb-secret\n"
    "REVERB_HOST=localhost\n"
    "REVERB_PORT=8080\n"
    "REVERB_SCHEME=http\n\n"
    "BCRYPT_ROUNDS=12"
)
gap()

# ============================================================
#  MENJALANKAN APLIKASI
# ============================================================
heading("Menjalankan Aplikasi", 1)
para("Jalankan masing-masing di terminal terpisah:")
gap()
run_steps = [
    ("1. Web server (development)",      "php artisan serve"),
    ("2. Queue worker",                  "php artisan queue:work --tries=3"),
    ("3. Reverb WebSocket server",       "php artisan reverb:start"),
    ("4. Scheduler (opsional — via cron)","* * * * * cd /path/project && php artisan schedule:run >> /dev/null 2>&1"),
]
for title, cmd in run_steps:
    para(title, bold=True)
    code_block(cmd)
    gap()

# ============================================================
#  DEFAULT LOGIN
# ============================================================
heading("Default Login", 1)
table(
    ["Field", "Value"],
    [["Username", "admin"], ["Password", "admin123"]],
)
gap()
para("Ganti password setelah login pertama kali.", italic=True)
gap()

# ============================================================
#  ARTISAN COMMANDS
# ============================================================
heading("Perintah Artisan Kustom", 1)
table(
    ["Perintah", "Keterangan"],
    [
        [
            "php artisan maintenance:backfill-signatures",
            "Generate/regenerasi Digital Signature + Payload Snapshot untuk semua "
            "maintenance log yang ada. Wajib dijalankan setelah pertama kali "
            "menambahkan fitur Digital Signature ke data yang sudah ada.",
        ],
        [
            "php artisan schedules:check-deadlines",
            "Cek semua jadwal: tandai yang overdue, broadcast alert untuk yang "
            "mendekati deadline (3 hari).",
        ],
    ],
)
gap()

# ============================================================
#  LISENSI
# ============================================================
heading("Lisensi", 1)
para("MIT License — bebas digunakan dan dimodifikasi.")
gap()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Built with Laravel 11  |  "
    "Dikembangkan untuk keperluan Ujian Akhir Semester (UAS) — Keamanan Data dan Informasi"
)
r.italic, r.font.size, r.font.name = True, Pt(10), "Calibri"

# ── Save ─────────────────────────────────────────────────────────────────────
output = "README_SIMM.docx"
doc.save(output)
print(f"SUCCESS: {output} berhasil dibuat!")
